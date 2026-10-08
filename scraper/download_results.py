import requests
import json
import os

API_BASE_URL = "https://loteriascaixa-api.herokuapp.com/api/lotofacil"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
}


def format_contest_item(item):
    """Padroniza o dicionário de um concurso."""
    if not item or not isinstance(item, dict):
        return None
    contest_num = item.get("concurso")
    date_str = item.get("data")
    dezenas = item.get("dezenas", [])
    if contest_num and dezenas:
        return {
            "concurso": int(contest_num),
            "data": date_str,
            "dezenas": sorted([int(d) for d in dezenas])
        }
    return None


def fetch_latest_contest():
    """Busca o concurso mais recente diretamente do endpoint /latest."""
    try:
        url = f"{API_BASE_URL}/latest"
        response = requests.get(url, headers=HEADERS, timeout=15)
        response.raise_for_status()
        return format_contest_item(response.json())
    except Exception as e:
        print(f"⚠️ Erro ao consultar último concurso: {e}")
        return None


def fetch_contest_by_number(contest_num):
    """Busca um concurso específico por número."""
    try:
        url = f"{API_BASE_URL}/{contest_num}"
        response = requests.get(url, headers=HEADERS, timeout=15)
        response.raise_for_status()
        return format_contest_item(response.json())
    except Exception as e:
        print(f"⚠️ Erro ao consultar concurso {contest_num}: {e}")
        return None


def fetch_all_results():
    """
    Busca todos os resultados da API, garantindo que o concurso mais recente
    (mesmo que não conste ainda no cache da lista completa) seja incluído.
    """
    try:
        response = requests.get(API_BASE_URL, headers=HEADERS, timeout=45)
        response.raise_for_status()
        data = response.json()
        
        if isinstance(data, list):
            contests = []
            for item in data:
                c = format_contest_item(item)
                if c:
                    contests.append(c)
            
            contests = sorted(contests, key=lambda x: x["concurso"])
            
            # Verifica se o endpoint /latest possui concurso mais novo do que o lote
            latest = fetch_latest_contest()
            if latest and contests and latest["concurso"] > contests[-1]["concurso"]:
                last_concurso = contests[-1]["concurso"]
                for num in range(last_concurso + 1, latest["concurso"] + 1):
                    if num == latest["concurso"]:
                        contests.append(latest)
                    else:
                        missing = fetch_contest_by_number(num)
                        if missing:
                            contests.append(missing)
                contests = sorted(contests, key=lambda x: x["concurso"])

            return contests
        else:
            print("❌ Formato de resposta da API não reconhecido.")
            return None
    except Exception as e:
        print(f"❌ Erro ao buscar dados completos: {e}")
        return None


def save_results_to_excel(contests, excel_path="lotofacil_resultados.xlsx"):
    """
    Exporta a lista completa de concursos para planilha Excel (.xlsx).
    Formato: Concurso | Data | Bola1 .. Bola15
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("⚠️ Pacote 'openpyxl' não instalado. Planilha Excel não gerada.")
        return False

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        headers = ['Concurso', 'Data'] + [f'Bola{i}' for i in range(1, 16)]
        ws.append(headers)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="800080", end_color="800080", fill_type="solid")
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for c in contests:
            row = [c["concurso"], c["data"]] + sorted(c["dezenas"])
            ws.append(row)

        wb.save(excel_path)
        return True
    except Exception as e:
        print(f"❌ Erro ao salvar arquivo Excel '{excel_path}': {e}")
        return False


def sync_contests(json_path="lotofacil_resultados.json", excel_path="lotofacil_resultados.xlsx"):
    """
    Sincronização rápida e inteligente:
    1. Se já temos arquivo JSON local, busca o /latest (leva < 1s).
    2. Se houver novos concursos, baixa apenas os faltantes e atualiza.
    3. Se não houver arquivo local, faz a carga completa.
    4. Salva JSON e Excel simultaneamente.
    Retorna a lista atualizada de concursos (ou None em caso de falha).
    """
    contests = []
    if os.path.exists(json_path):
        try:
            with open(json_path, 'r', encoding='utf-8') as f:
                contests = json.load(f)
            contests = sorted(contests, key=lambda x: x["concurso"])
        except Exception as e:
            print(f"⚠️ Erro ao ler '{json_path}': {e}. Realizando carga completa...")
            contests = []

    # Consulta último concurso online
    latest = fetch_latest_contest()
    
    if contests and latest:
        last_local_num = contests[-1]["concurso"]
        latest_online_num = latest["concurso"]
        
        if latest_online_num > last_local_num:
            print(f"🔄 Baixando novos sorteios (do concurso {last_local_num + 1} ao {latest_online_num})...")
            for num in range(last_local_num + 1, latest_online_num + 1):
                if num == latest_online_num:
                    contests.append(latest)
                else:
                    item = fetch_contest_by_number(num)
                    if item:
                        contests.append(item)
            contests = sorted(contests, key=lambda x: x["concurso"])
        else:
            print(f"ℹ️ Base já contém o concurso mais recente (#{last_local_num}).")
    elif not contests:
        print("Iniciando download completo da base...")
        contests = fetch_all_results()

    if contests and len(contests) > 0:
        # Salva JSON
        with open(json_path, 'w', encoding='utf-8') as f:
            json.dump(contests, f, ensure_ascii=False, indent=2)
        
        # Salva Excel
        save_results_to_excel(contests, excel_path)
        return contests
    
    return None


if __name__ == "__main__":
    sync_contests()
